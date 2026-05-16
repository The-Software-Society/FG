"""Seed inputs.yaml from the original Excel financial model.

Run once after install. Idempotent — re-running overwrites inputs.yaml with
freshly-extracted values. Source spreadsheet path can be overridden with
`--xlsx`; defaults to the file in Jon's Downloads folder.

Usage:
    python scripts/extract_baseline_from_xlsx.py
    python scripts/extract_baseline_from_xlsx.py --xlsx /path/to/model.xlsx --out inputs.yaml
"""
from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import yaml

DEFAULT_XLSX = Path("/Users/jon/Downloads/The Software Society - Financial Model (1).xlsx")
DEFAULT_OUT = Path(__file__).resolve().parent.parent / "inputs.yaml"

SCENARIOS = ["conservative", "standard", "optimistic"]
TIERS = ["basic", "pro"]
HORIZON = 36


def _v(ws, coord):
    return ws[coord].value


def _vfloat(ws, coord, default=None):
    v = ws[coord].value
    if v is None:
        return default
    return float(v)


def _row_36(ws, col_letter, start_row=10):
    """Read a 36-row vertical strip starting at <col><start_row>."""
    return [_vfloat(ws, f"{col_letter}{start_row + i}", 0.0) for i in range(HORIZON)]


def _row_36_horizontal(ws, row, start_col_idx=8):
    """Read a 36-col horizontal strip starting at column index `start_col_idx` (1-based)."""
    out = []
    for i in range(HORIZON):
        c_idx = start_col_idx + i
        col = openpyxl.utils.get_column_letter(c_idx)
        v = _vfloat(ws, f"{col}{row}", 0.0)
        out.append(v)
    return out


def _step_schedule(ws, start_row, end_row):
    """Read a step table with cols B (start_month), C (final_month), D (monthly).

    Excel rows where B is "NA" or empty are skipped. Returns list of dicts.
    """
    sched = []
    for r in range(start_row, end_row + 1):
        b = _v(ws, f"B{r}")
        c = _v(ws, f"C{r}")
        d = _v(ws, f"D{r}")
        if b is None or b == "NA" or c is None or c == "NA":
            continue
        try:
            sched.append({
                "months": [int(b), int(c)],
                "monthly": float(d if d is not None else 0.0),
            })
        except (TypeError, ValueError):
            continue
    return sched


def _scenario_dict(values_per_scenario):
    """Given (cons, std, opt), return dict keyed by scenario."""
    return dict(zip(SCENARIOS, [float(x) if x is not None else 0.0 for x in values_per_scenario]))


def _tier_scenario_block(ws, basic_row, pro_row, cols=("C", "D", "E")):
    """Read a 2-row × 3-col block (basic/pro × cons/std/opt)."""
    return {
        "basic": _scenario_dict([_v(ws, f"{col}{basic_row}") for col in cols]),
        "pro":   _scenario_dict([_v(ws, f"{col}{pro_row}")   for col in cols]),
    }


def extract(xlsx_path: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    a = wb["Assumptions"]
    s = wb["Statements"]

    # ── Revenue stream ──────────────────────────────────────────────────
    new_clients = {
        "conservative": _row_36(a, "C", 10),
        "standard":     _row_36(a, "D", 10),
        "optimistic":   _row_36(a, "E", 10),
    }

    revenue_stream = {
        "name": _v(a, "C6") or "Revenue Stream 1",
        "business_model": _v(a, "C7") or "",
        "new_clients_per_month": new_clients,
        "churn_monthly": _vfloat(a, "C73", 0.0),
        "tier_split": {
            "conservative": {"basic": _vfloat(a, "C48"), "pro": _vfloat(a, "C49")},
            "standard":     {"basic": _vfloat(a, "D48"), "pro": _vfloat(a, "D49")},
            "optimistic":   {"basic": _vfloat(a, "E48"), "pro": _vfloat(a, "E49")},
        },
        "pricing": {
            "subscription": _tier_scenario_block(a, 53, 54),
            "seats": {
                "per_client": _tier_scenario_block(a, 57, 58),
                "price":      _tier_scenario_block(a, 61, 62),
            },
            "credits": {
                "per_client_monthly": _tier_scenario_block(a, 65, 66),
                "price":              _tier_scenario_block(a, 69, 70),
            },
        },
        "invoice_fulfillment_pct": _vfloat(a, "C76", 1.0),
        "collection_lag_months":   int(_vfloat(a, "C77", 0)),
    }

    # ── Cost streams ────────────────────────────────────────────────────
    cost_streams = [
        {
            "name": _v(a, "B87") or "Cost Stream 1",
            "kind": "step",
            "schedule": _step_schedule(a, 89, 94),
        },
        {
            "name": _v(a, "B96") or "Cost Stream 2",
            "kind": "per_active_client",
            "per_active_client_by_year": [
                _vfloat(a, "C97", 0.0),
                _vfloat(a, "D97", 0.0),
                _vfloat(a, "E97", 0.0),
            ],
        },
        {
            "name": _v(a, "B99") or "Cost Stream 3",
            "kind": "per_credit",
            "per_credit_by_year": [
                _vfloat(a, "C100", 0.0),
                _vfloat(a, "D100", 0.0),
                _vfloat(a, "E100", 0.0),
            ],
        },
    ]

    # ── Salaries ────────────────────────────────────────────────────────
    role_rows = [
        ("Management: Co-Founders",         109),
        ("Management: C-Level",             110),
        ("Management: High-Level Associates", 111),
        ("Associates: Legal Support",       112),
        ("Associates: Finance & Accounting", 113),
        ("Associates: HR & People",         114),
        ("Associates: Tech",                115),
        ("Associates: Marketing",           116),
        ("Associates: Operations",          117),
    ]
    salaries_roles = []
    for name, row in role_rows:
        base = _vfloat(a, f"C{row}", 0.0) or 0.0
        # Headcount in cols H..AQ (idx 8..43) of `row`
        headcount = [int(x) for x in _row_36_horizontal(a, row, start_col_idx=8)]
        salaries_roles.append({
            "name": name,
            "base_monthly": base,
            "headcount_by_month": headcount,
        })

    salaries = {
        "fringe_benefits_pct":         _vfloat(a, "C120", 0.0),
        "payment_periodicity_months":  int(_vfloat(a, "C121", 1)),
        "roles":                       salaries_roles,
    }

    # ── OpEx (9 step tables) ────────────────────────────────────────────
    opex_specs = [
        ("Travel Expenses",          126, 131),
        ("Office Space Expenses",    136, 141),
        ("Office Supplies Expenses", 146, 151),
        ("Tech Licenses Expenses",   156, 161),
        ("Banking Expenses",         166, 171),
        ("Advisory Expenses",        176, 181),
        ("Marketing Expenses",       186, 191),
        ("Customer Support Expenses", 196, 201),
        ("Unforeseen/Miscellaneous Expenses", 206, 211),
    ]
    opex = [
        {"name": name, "schedule": _step_schedule(a, sr, er)}
        for name, sr, er in opex_specs
    ]

    # ── CAPEX & Tech (3 slots each) ────────────────────────────────────
    def _capex_slot(value_row, life_row, acq_row, name):
        v = _vfloat(a, f"C{value_row}")
        L = _vfloat(a, f"C{life_row}")
        m = _vfloat(a, f"C{acq_row}")
        if v is None or L is None or m is None:
            return None
        return {
            "name": name,
            "value": v,
            "depreciation_months": int(L),
            "acquisition_month": int(m),
        }

    capex_raw = [
        _capex_slot(219, 220, 221, "CAPEX Investment #1"),
        _capex_slot(225, 226, 227, "CAPEX Investment #2"),
        _capex_slot(231, 232, 233, "CAPEX Investment #3"),
    ]
    capex = [c for c in capex_raw if c is not None]

    def _tech_slot(value_row, life_row, acq_row, name):
        v = _vfloat(a, f"C{value_row}")
        L = _vfloat(a, f"C{life_row}")
        m = _vfloat(a, f"C{acq_row}")
        if v is None or L is None or m is None:
            return None
        return {
            "name": name,
            "value": v,
            "amortization_months": int(L),
            "acquisition_month": int(m),
        }

    tech_raw = [
        _tech_slot(241, 242, 243, "Tech Investment #1"),
        _tech_slot(247, 248, 249, "Tech Investment #2"),
        _tech_slot(253, 254, 255, "Tech Investment #3"),
    ]
    tech = [t for t in tech_raw if t is not None]

    # ── Debt (3 slots) ──────────────────────────────────────────────────
    def _debt_slot(prn_row, term_row, rate_row, dis_row, name):
        p = _vfloat(a, f"C{prn_row}")
        t = _vfloat(a, f"C{term_row}")
        r = _vfloat(a, f"C{rate_row}")
        d = _vfloat(a, f"C{dis_row}")
        if p is None or t is None or r is None or d is None:
            return None
        return {
            "name": name,
            "principal": p,
            "term_months": int(t),
            "monthly_rate": r,
            "disbursement_month": int(d),
        }

    debt_raw = [
        _debt_slot(263, 264, 265, 266, "Debt #1"),
        _debt_slot(270, 271, 272, 273, "Debt #2"),
        _debt_slot(277, 278, 279, 280, "Debt #3"),
    ]
    debt = [d for d in debt_raw if d is not None]

    # ── Equity raises (3 slots) ─────────────────────────────────────────
    def _equity_slot(amt_row, mo_row, name):
        amt = _vfloat(a, f"C{amt_row}")
        mo  = _vfloat(a, f"C{mo_row}")
        if amt is None or mo is None:
            return None
        return {"name": name, "amount": amt, "closing_month": int(mo)}

    equity_raises_raw = [
        _equity_slot(288, 289, "Fundraise #1"),
        _equity_slot(293, 294, "Fundraise #2"),
        _equity_slot(298, 299, "Fundraise #3"),
    ]
    equity_raises = [e for e in equity_raises_raw if e is not None]

    # ── Grants (3 slots) ────────────────────────────────────────────────
    grants_raw = [
        _equity_slot(307, 308, "Grant #1"),
        _equity_slot(312, 313, "Grant #2"),
        _equity_slot(317, 318, "Grant #3"),
    ]
    grants = [{"name": g["name"], "amount": g["amount"], "receiving_month": g["closing_month"]}
              for g in grants_raw if g is not None]

    # ── Other Income ────────────────────────────────────────────────────
    other_income = {"schedule": _step_schedule(a, 327, 332)}

    # ── Tax ─────────────────────────────────────────────────────────────
    tax = {"effective_rate": _vfloat(s, "B21", 0.21)}

    return {
        "horizon_months": HORIZON,
        "scenarios": SCENARIOS,
        "revenue_streams": [revenue_stream],
        "payment_lag_months": int(_vfloat(a, "C85", 1)),
        "cost_streams": cost_streams,
        "salaries": salaries,
        "opex": opex,
        "capex": capex,
        "tech_investments": tech,
        "debt": debt,
        "equity_raises": equity_raises,
        "grants": grants,
        "other_income": other_income,
        "tax": tax,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--out",  type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    if not args.xlsx.exists():
        raise SystemExit(f"Source xlsx not found: {args.xlsx}")

    data = extract(args.xlsx)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w") as f:
        yaml.dump(
            data, f, sort_keys=False, allow_unicode=True, default_flow_style=False, width=120
        )

    print(f"Wrote {args.out}")
    print(f"  - {len(data['revenue_streams'])} revenue stream(s)")
    print(f"  - {len(data['cost_streams'])} cost stream(s)")
    print(f"  - {len(data['salaries']['roles'])} salary role(s)")
    print(f"  - {len(data['opex'])} opex categor(ies)")
    print(f"  - {len(data['capex'])} CAPEX, {len(data['tech_investments'])} tech, "
          f"{len(data['debt'])} debt, {len(data['equity_raises'])} equity, "
          f"{len(data['grants'])} grant(s)")


if __name__ == "__main__":
    main()
